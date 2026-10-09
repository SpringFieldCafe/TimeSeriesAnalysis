"""Fill the local experiment report template from analysis.py results."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent


def scientific(value: float) -> str:
    if value >= 0.001:
        return f"{value:.4f}"
    mantissa, exponent = f"{value:.2e}".split("e")
    superscript = str(int(exponent)).translate(str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹"))
    return f"{mantissa}×10{superscript}"


def format_run(run, size=12, bold=False, font="宋体"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def disable_grid(p):
    element = OxmlElement("w:snapToGrid")
    element.set(qn("w:val"), "0")
    p._p.get_or_add_pPr().append(element)


def math_run(text):
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_container(name, *children):
    element = OxmlElement(name)
    for child in children:
        element.append(child)
    return element


def subscript(text, sub):
    return math_container("m:sSub", math_container("m:e", math_run(text)), math_container("m:sub", math_run(sub)))


def add_equation(p, kind):
    equation = OxmlElement("m:oMath")
    if kind == "return":
        equation.append(subscript("r", "t"))
        equation.append(math_run(" = 100 ln "))
        equation.append(math_container("m:f", math_container("m:num", subscript("P", "t")), math_container("m:den", subscript("P", "t−1"))))
    else:
        equation.append(math_run("Q(m) = n(n+2) "))
        nary = OxmlElement("m:nary")
        props = OxmlElement("m:naryPr")
        char = OxmlElement("m:chr")
        char.set(qn("m:val"), "∑")
        location = OxmlElement("m:limLoc")
        location.set(qn("m:val"), "subSup")
        props.extend([char, location])
        nary.append(props)
        nary.append(math_container("m:sub", math_run("k=1")))
        nary.append(math_container("m:sup", math_run("m")))
        squared = math_container("m:sSup", math_container("m:e", subscript("r", "k")), math_container("m:sup", math_run("2")))
        fraction = math_container("m:f", math_container("m:num", squared), math_container("m:den", math_run("n−k")))
        nary.append(math_container("m:e", fraction))
        equation.append(nary)
    p._p.append(equation)


def fill_report(template: Path, output: Path):
    data = json.loads((ROOT / "results.json").read_text(encoding="utf-8"))
    a, c, s, lb = data["apple"], data["co2"], data["sales"], data["sales_ljung_box"]
    doc = Document(template)
    original = list(doc.paragraphs)
    if len(original) < 70 or original[46].text.strip() != "1.":
        raise ValueError("Use the original unfilled report template, not a completed report")
    prototype = deepcopy(original[46]._p)
    anchor = original[54]._p
    for p in original[46:54]:
        p._p.getparent().remove(p._p)

    def paragraph(text="", *, heading=False, new_page=False, center=False, size=12, indent=True):
        element = deepcopy(prototype)
        p = Paragraph(element, doc._body)
        p.clear()
        fmt = p.paragraph_format
        fmt.space_before = Pt(8 if heading else 0)
        if anchor.getprevious().tag == qn("w:tbl"):
            fmt.space_before = Pt(6)
        fmt.space_after = Pt(5 if heading else 4)
        fmt.line_spacing = 1.2
        fmt.keep_with_next = heading
        fmt.widow_control = True
        fmt.page_break_before = new_page
        fmt.first_line_indent = Cm(0.85 if indent and not heading and not center else 0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        disable_grid(p)
        format_run(p.add_run(text), size=size, bold=heading)
        anchor.addprevious(element)
        return p

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
        for cells, values in [(t.rows[0].cells, headers)] + [(t.add_row().cells, row) for row in rows]:
            for index, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Cm(widths[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                disable_grid(p)
                format_run(p.add_run(str(value)), size=10.5, bold=cells is t.rows[0].cells)
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ["top", "left", "bottom", "right"]:
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "45" if side in ["top", "bottom"] else "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
        borders = OxmlElement("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{side}")
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        t._tbl.tblPr.append(borders)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell in t.rows[0].cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EDEDED")
            cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        anchor.addprevious(t._tbl)

    def figure(filename, caption, width=13.2):
        p = paragraph(center=True, indent=False)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(1)
        p.add_run().add_picture(str(ROOT / "figures" / filename), width=Cm(width))
        picture = p._p.xpath('.//wp:docPr')[0]
        picture.set("descr", caption)
        cp = paragraph(caption, center=True, size=10.5, indent=False)
        cp.paragraph_format.space_after = Pt(6)

    paragraph("1 数据读取与实验环境", heading=True)
    paragraph("在本机 Anaconda 中运行 Python，使用 pandas 读取数据、matplotlib 绘图、statsmodels 完成统计检验。")
    table(["序列", "实际起止月份", "样本数", "缺失值"], [
        ["Apple 股价", "2020-01—2025-12", a["count"], a["missing"]],
        ["CO₂ 月度数据", "1975-01—1980-12", c["count"], c["missing"]],
        ["月度销售量", "2000-01—2003-12", s["count"], s["missing"]],
    ], [3.2, 6, 2.6, 2.6])
    paragraph("Apple 采用美元计价的拆股与分红复权收盘价，非交易日不补值。数据保存在 doc/，图表保存在 figures/。月度数据来自实验资料。")
    source = json.loads((ROOT / "doc" / "AAPL_2020_2025_source.json").read_text(encoding="utf-8"))
    p = paragraph(f"股价数据下载日期 {source['retrieved_at_utc'][:10]}，来源 ", size=10.5, indent=False)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), doc.part.relate_to(source["source_page"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "Yahoo Finance AAPL 历史行情"
    run.append(text)
    link.append(run)
    p._p.append(link)
    paragraph("2 分析方法与判断依据", heading=True)
    paragraph("图检验观察趋势、均值与季节波动，并结合 ACF 判断；股价滞后单位为交易日，月度数据为月。Apple 计算 20 日和 60 日移动平均及对数收益率。ADF 含常数项，以 AIC 选滞后阶数，原假设为存在单位根；LB 检验取 10 阶。显著性水平均为 0.05。")

    paragraph("3 Apple 公司股价探索", heading=True, new_page=True)
    paragraph(f"选取 2020—2025 年日度数据，实际范围为 {a['first_date']} 至 {a['last_date']}，共 {a['count']} 个交易日。复权收盘价从 {a['first_price']:.2f} 美元变为 {a['last_price']:.2f} 美元；区间最低为 {a['min']:.2f} 美元（{a['min_date']}），最高为 {a['max']:.2f} 美元（{a['max_date']}）。")
    figure("apple_price_time.png", "图 1 Apple 复权收盘价及 20 日与 60 日移动平均")
    paragraph("图 1 显示股价整体水平随时间上升，期间存在回落，移动平均线也随时间变化。价格序列围绕固定均值波动的特征不明显，按图检验判断为非平稳。移动平均仅用于描述趋势。")
    figure("apple_price_acf.png", "图 2 Apple 复权收盘价样本自相关图")
    paragraph(f"价格的一阶自相关为 {a['acf'][1]:.3f}，24 阶仍为 {a['acf'][24]:.3f}，呈缓慢衰减。ADF 统计量为 {a['price_adf']['statistic']:.3f}，p 值为 {a['price_adf']['p_value']:.4f}，在 5% 水平下未拒绝单位根假设，与图形所提示的非平稳特征一致。")

    paragraph("Apple 对数收益率与相关结构", heading=True, new_page=True)
    p = paragraph(center=True, indent=False)
    add_equation(p, "return")
    paragraph(f"式中 P 为复权收盘价，r 为百分数形式的日对数收益率。差分后的有效样本数为 {a['return_count']}，均值为 {a['return_mean_pct']:.4f}%，样本标准差为 {a['return_std_pct']:.4f}%。")
    figure("apple_returns_time.png", "图 3 Apple 日对数收益率时序图", width=11.2)
    figure("apple_returns_acf.png", "图 4 Apple 日对数收益率样本自相关图", width=11.2)
    table(["检验对象与方法", "统计量", "p 值"], [
        ["收益率 ADF", f"{a['return_adf']['statistic']:.3f}", scientific(a['return_adf']['p_value'])],
        ["收益率 LB(10)", f"{a['return_ljung_box']['statistic']:.3f}", scientific(a['return_ljung_box']['p_value'])],
        ["平方收益率 LB(10)", f"{a['squared_return_ljung_box']['statistic']:.3f}", scientific(a['squared_return_ljung_box']['p_value'])],
    ], [7, 3.4, 4])
    paragraph("收益率围绕零附近波动，ADF 拒绝单位根假设，但 LB(10) 拒绝零自相关假设。平方收益率存在显著相关性，与图中的波动聚集相符，不能把收益率视为独立白噪声；ADF 也不能单独证明方差稳定。")

    paragraph("4 CO₂ 月度序列的图检验", heading=True, new_page=True)
    paragraph(f"读取 E2_2.xlsx 的 time 与 co2 列。共 {c['count']} 个观测，均值 {c['mean']:.3f}，样本标准差 {c['std']:.3f}，最小值 {c['min']:.2f}，最大值 {c['max']:.2f}。")
    figure("co2_time.png", "图 5 1975—1980 年 CO₂ 月度时序图")
    paragraph(f"浓度具有持续上升趋势，同时每年重复出现峰谷。年度均值从 1975 年的 {c['annual_mean']['1975']:.3f} 上升至 1980 年的 {c['annual_mean']['1980']:.3f}，说明均值水平随时间改变。按图检验，该序列不平稳，季节周期约为 12 个月。")
    figure("co2_acf.png", "图 6 CO₂ 月度序列样本自相关图")
    paragraph(f"一阶自相关为 {c['acf'][1]:.3f}，随后缓慢下降，并在 12 阶附近回升至 {c['acf'][12]:.3f}，24 阶约为 {c['acf'][24]:.3f}。趋势和年度季节性共同形成持续相关及周期性起伏，进一步支持图检验的非平稳判断。")

    paragraph("5 月度销售量的图检验", heading=True, new_page=True)
    paragraph(f"读取 E2_5.xlsx 的 time 与 x 列。共 {s['count']} 个观测，均值为 {s['mean']:.3f}，样本标准差为 {s['std']:.3f}，取值范围为 {s['min']:.0f}—{s['max']:.0f}，单位沿用实验数据。")
    figure("sales_time.png", "图 7 2000—2003 年月度销售量时序图")
    paragraph(f"销售量多在春末初夏达到高值，年末处于低值，季节波动明显。四年的年均值依次为 {s['annual_mean']['2000']:.3f}、{s['annual_mean']['2001']:.3f}、{s['annual_mean']['2002']:.3f} 和 {s['annual_mean']['2003']:.3f}，后期水平下降。按图检验，原序列不平稳。")
    figure("sales_acf.png", "图 8 月度销售量样本自相关图")
    paragraph(f"一阶自相关约为 {s['acf'][1]:.3f}，6 阶为 {s['acf'][6]:.3f}，12 阶为 {s['acf'][12]:.3f}，24 阶为 {s['acf'][24]:.3f}。正负相关交替，约 12 个月重复一次，与时序图显示的年内季节性一致。季节性均值变化是非平稳判断的主要依据。")

    paragraph("6 月度销售量的纯随机性检验", heading=True, new_page=True)
    paragraph("采用 Ljung–Box 检验。原假设 H₀ 为前 10 阶总体自相关系数均等于零，备择假设 H₁ 为至少一阶不为零。检验显著性水平取 0.05，统计量为")
    p = paragraph(center=True, indent=False)
    add_equation(p, "lb")
    paragraph("n 为样本量，m 为检验滞后阶数，rₖ 为 k 阶样本自相关。这里 n=48、m=10，未拟合 ARMA 模型，使用 10 个自由度的卡方分布近似计算 p 值。")
    table(["样本数 n", "滞后阶数 m", "Q 统计量", "p 值"], [
        [s["count"], lb["lag"], f"{lb['statistic']:.6f}", scientific(lb["p_value"])],
    ], [2.5, 3.1, 4.2, 4.6])
    paragraph("主要计算代码如下，完整分析程序见 exp1/analysis.py。", indent=False)
    code = [
        'sales = load_series("E2_5.xlsx", "x")',
        'result = acorr_ljungbox(',
        '    sales, lags=[10], model_df=0, return_df=True)',
        'print(result)',
    ]
    for line in code:
        p = paragraph(line, size=9.5, indent=False)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            format_run(run, size=9.5, font="Consolas")
    paragraph(f"计算得 Q(10)={lb['statistic']:.3f}，p={scientific(lb['p_value'])}，远小于 0.05，拒绝原假设。结合显著的季节性自相关，销售量原序列不是纯随机序列。该检验用于检验自相关，不能代替平稳性检验。")

    summaries = [
        "Apple 价格表现出趋势和持续自相关，对数收益率的均值更稳定，但仍有相关性与波动聚集。CO₂ 具有上升趋势和年度季节性；销售量具有年内波动且后期下降，两者按图检验均不平稳。",
        "销售量 LB 检验拒绝零自相关假设，原序列不是纯随机序列。平稳性与纯随机性回答不同问题。后续可考虑趋势差分与季节差分，再重新检验；本实验的月度数据按题目要求进行原序列图检验。",
        "",
    ]
    for p, text in zip(original[55:58], summaries):
        p.clear()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0.85)
        disable_grid(p)
        format_run(p.add_run(text))
    # Keep the instructor fields together, retaining their original text and fonts.
    for p in original[60:64] + [original[65], original[67], original[69]]:
        p._p.getparent().remove(p._p)
    for p in [original[58], original[59], original[64], original[66]]:
        p.paragraph_format.keep_with_next = True

    buffer = BytesIO()
    doc.save(buffer)
    # Merge only authored package parts; all other original parts stay byte-identical.
    editable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(template) as source, ZipFile(buffer) as authored, ZipFile(output, "w") as final:
        for info in source.infolist():
            final.writestr(info, authored.read(info.filename) if info.filename in editable else source.read(info.filename))
        for info in authored.infolist():
            if info.filename not in source.namelist():
                final.writestr(info, authored.read(info.filename))
    print(f"Report saved to: {output}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "report" / "第一个实验报告.docx")
    args = parser.parse_args()
    template = args.template
    if template is None:
        template = ROOT / "report" / "报告模板.docx"
        if not template.exists():
            shutil.copy2(args.output, template)
    fill_report(template, args.output)


if __name__ == "__main__":
    main()
