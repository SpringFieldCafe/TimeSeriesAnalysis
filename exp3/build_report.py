"""从实验三计算结果生成本地 Word 报告。先运行 analysis.py。"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent


def main():
    results = json.loads((ROOT / "results" / "results.json").read_text(encoding="utf-8"))
    a, b, diff = [results[k] for k in ("e2_7", "e2_8", "e2_8_diff")]
    doc = Document()
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "宋体"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].paragraph_format.space_after = Pt(5)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "时间序列分析实验三"
    doc.core_properties.subject = "平稳性与自相关和偏自相关特征"
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""

    def paragraph(text):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control = True
        return p

    def heading(text, new_page=False):
        p = doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = new_page
        return p

    def image(name, caption):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ROOT / "figures" / name), width=Cm(16.4))
        p = doc.add_paragraph(caption, style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for cells, values in [(t.rows[0].cells, headers)] + [(t.add_row().cells, row) for row in rows]:
            for cell, value, width in zip(cells, values, widths):
                cell.width = Cm(width)
                cell.text = str(value)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for run in p.runs:
                        run.font.size = Pt(10)
        for cell in t.rows[0].cells:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EAEAEA")
            cell._tc.get_or_add_tcPr().append(shading)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        header = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(header)
        for row in t.rows:
            element = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(element)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph("时间序列分析实验三", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph("本实验通过时序图、样本自相关函数（ACF）和偏自相关函数（PACF），分析澳大利亚枪支相关凶杀案死亡率及密歇根湖水位的平稳特征。两组原序列的 ACF 均呈拖尾，PACF 可近似描述为二阶截尾。澳大利亚死亡率的平稳性证据不一致；密歇根湖原水位按非平稳序列处理，一阶差分支持近似平稳。")
    heading("数据与实验要求")
    table(["数据", "观测范围", "样本数", "变量含义"], [
        ["E2_7", "1915–2004年", a["n"], "死亡率 每10万人"],
        ["E2_8", "1860–1955年", b["n"], "每年月平均水位最高值"],
    ], [2.1, 4.2, 2, 8.2])
    paragraph("两表均为连续年度观测，无缺失值。水位表未注明单位，保留原数据单位；题目中的月平均水位最高值每年记录一次。各题均要求绘制时序图、考察平稳性，并判断 ACF 与 PACF 的截尾或拖尾特征。")
    heading("分析方法")
    paragraph("时序图同时绘制10年居中移动平均以观察局部均值。ACF/PACF 计算1–24阶；ACF 使用 Bartlett 95%逐阶参考区间，PACF 使用 Yule–Walker 估计（ywm），参考半宽约为1.96/√n。非平稳序列的区间仅作图形参考，移动平均不参与检验。")
    paragraph("ADF 以存在单位根为原假设，使用常数项与 AIC 选滞后阶数；KPSS 以水平平稳为原假设，使用常数项与自动带宽。两项检验均采用5%显著性水平。KPSS 边界 p 值用不等号表示；未拒绝原假设不等于证明它成立。")
    table(["序列", "ADF统计量", "ADF p值", "KPSS统计量", "KPSS p值"], [
        [label, f"{r['adf_stat']:.4f}", f"{r['adf_p']:.4f}" if r['adf_p'] >= .001 else f"{r['adf_p']:.2e}", f"{r['kpss_stat']:.4f}", f"{r['kpss_p_bound']}{r['kpss_p']:.2f}"]
        for label, r in [("死亡率", a), ("原水位", b), ("水位一阶差分", diff)]
    ], [4, 3.1, 3.1, 3.1, 3.2])
    paragraph("截尾和拖尾是理论相关函数的形态。平稳 AR(p) 的 ACF 拖尾、PACF 截尾；平稳 MA(q) 的 ACF 截尾、PACF 拖尾。样本图只能作近似判断，逐阶区间未作多重比较校正。原序列若非平稳，应先平稳化再进行模型识别。")

    for prefix, label, r in [("e2_7", "澳大利亚死亡率分析", a), ("e2_8", "密歇根湖水位分析", b)]:
        heading(label, new_page=True)
        paragraph(f"观测期 {r['start_year']}–{r['end_year']} 年，n={r['n']}。均值 {r['mean']:.4f}，样本标准差 {r['sd']:.4f}，最小值 {r['min']:.4f}，最大值 {r['max']:.4f}。")
        image(f"{prefix}_time.png", "时序图与10年居中移动平均")
        paragraph(r["stationarity"])
        image(f"{prefix}_acf_pacf.png", "样本 ACF 与 PACF 及95%逐阶参考区间")
        paragraph(r["correlation"])
        c = r["correlations"]
        paragraph(f"PACF 前三阶为 {c[1]['pacf']:.4f}、{c[2]['pacf']:.4f}、{c[3]['pacf']:.4f}，参考半宽为 {r['pacf_reference_halfwidth']:.4f}。ACF 第1、5、10阶为 {c[1]['acf']:.4f}、{c[5]['acf']:.4f}、{c[10]['acf']:.4f}。")

    heading("水位一阶差分与实验结论", new_page=True)
    paragraph("对密歇根湖水位作补充验证：Δx_t = x_t − x_(t−1)，保留后一年的年份作为差分索引，差分后有95个观测值。")
    image("e2_8_diff_time.png", "水位一阶差分时序图")
    paragraph(diff["stationarity"])
    image("e2_8_diff_acf_pacf.png", "水位一阶差分 ACF 与 PACF")
    paragraph(diff["correlation"])
    doc.add_paragraph("实验结论", style="Heading 2")
    paragraph("澳大利亚死亡率的 ACF 拖尾，PACF 近似二阶截尾，但第二阶临界、平稳性证据不一致，应保留判断。密歇根湖原水位的 ACF 缓慢拖尾，PACF 主要集中在前两阶；原序列非平稳，一阶差分后支持近似平稳。不能从这些样本图直接断言两者都服从 AR(2)。")
    reference = paragraph("方法参考：statsmodels 官方文档中 acf、pacf、adfuller 与 kpss 的定义与区间说明。https://www.statsmodels.org/stable/")
    reference.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.runs[0].font.size = Pt(9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = ROOT / "report" / "时间序列分析实验3报告.docx"
    output.parent.mkdir(exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
